#!/usr/bin/env python3
"""Omsi-Hub's texts as one spreadsheet for a translator, and the spreadsheet back into the tables
(crates/hub-locales/locales/*.yml).

    python tools/translations.py export it                 -> omsihub-it.xlsx
    python tools/translations.py import it omsihub-it.xlsx

The export has every English text once (a text in several tables is one row), the translations
into the reference languages beside it (Spanish and French unless --ref says other ones) and a
column for the language, filled where the tables have it already. The import writes each filled
cell into every table that has its English text: an `  it: "…"` line after the entry's last
translation, or in place of the one there. A text whose %{placeholders} differ from the English
one's is reported and left out. Needs openpyxl.
"""

import argparse
import glob
import os
import re
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

LOCALES = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "crates", "hub-locales", "locales")
NAMES = {"de": "Deutsch", "es": "Español", "fr": "Français", "it": "Italiano", "nl": "Nederlands", "pl": "Polski",
         "pt": "Português", "ru": "Русский", "uk": "Українська", "cs": "Čeština", "tr": "Türkçe"}
PLACEHOLDER = re.compile(r"%\{\w+\}")


def unquote(s):
    """A table's text as hub-locales reads it (`unquote` in its lib.rs)."""
    s = s.strip()
    if not (len(s) >= 2 and s[0] == '"' and s[-1] == '"'):
        return s
    out, it = [], iter(s[1:-1])
    for c in it:
        if c != "\\":
            out.append(c)
            continue
        n = next(it, None)
        if n is not None:
            out.append({"n": "\n", "t": "\t"}.get(n, n))
    return "".join(out)


def quote(s):
    return '"' + s.replace("\\", "\\\\").replace('"', '\\"').replace("\n", "\\n").replace("\t", "\\t") + '"'


def tables():
    """Every table: its path, its lines (without their ends) and its line end."""
    for path in sorted(glob.glob(os.path.join(LOCALES, "*.yml"))):
        with open(path, encoding="utf-8", newline="") as f:
            text = f.read()
        eol = "\r\n" if "\r\n" in text else "\n"
        yield path, text.split(eol), eol


def entries(lines):
    """The entries of a table's lines: (English text, line of its key, {lang: (line, text)})."""
    out = []
    for i, line in enumerate(lines):
        if not line.strip() or line.lstrip().startswith("#") or line.startswith("_version"):
            continue
        if line.startswith("  "):
            lang, _, text = line[2:].partition(":")
            if out:
                out[-1][2][lang.strip()] = (i, unquote(text))
        elif line.rstrip().endswith(":"):
            out.append((unquote(line.rstrip()[:-1]), i, {}))
    return out


def export(lang, refs, path):
    rows = {}
    for table, lines, _ in tables():
        name = os.path.basename(table)[:-4]
        for key, _, langs in entries(lines):
            row = rows.setdefault(key, {"tables": [], "langs": {}})
            row["tables"].append(name)
            for l, (_, text) in langs.items():
                row["langs"].setdefault(l, text)

    wb = Workbook()
    ws = wb.active
    ws.title = "Texts"
    head = ["#", "English"] + [NAMES.get(r, r) for r in refs] + [NAMES.get(lang, lang) + " (" + lang + ")", "Where"]
    ws.append(head)
    for n, (key, row) in enumerate(rows.items(), 1):
        ws.append([n, key] + [row["langs"].get(r, "") for r in refs] + [row["langs"].get(lang, ""), ", ".join(row["tables"])])
    # (a text starting with "=" stays text, not a formula)
    for cells in ws.iter_rows(min_row=2):
        for c in cells:
            if isinstance(c.value, str) and c.value.startswith("="):
                c.data_type = "s"
    col = len(head) - 1
    fill = PatternFill("solid", fgColor="FFF4CC")
    for c in ws[1]:
        c.font = Font(bold=True)
    for cells in ws.iter_rows(min_row=1, min_col=col, max_col=col):
        cells[0].fill = fill
    for i, width in enumerate([7, 60] + [45] * len(refs) + [60, 24], 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = width
    for cells in ws.iter_rows(min_row=2):
        for c in cells:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions

    how = wb.create_sheet("How to")
    for line in [
        "Translating Omsi-Hub",
        "",
        "Fill the yellow column. Leave the other columns as they are: the English text is how a translation finds its place.",
        "The other languages are there to help: the English text is the original.",
        "Keep %{name} placeholders exactly as they are (they are filled with a number, a name, ...). Their order may change.",
        "Keep line breaks, and keep names such as Omsi-Hub, RosterFlow, OMSI, IBIS and openOMSI.",
        "A cell left empty shows the English text in the app, so any part translated is a help.",
        "'Where' says which part of the app a text belongs to (company = the bus company, rosterflow = the planning, ...).",
        "Filter or sort as you like; save as .xlsx (Excel, LibreOffice or Google Sheets: File > Download > .xlsx).",
    ]:
        how.append([line])
    how["A1"].font = Font(bold=True, size=14)
    how.column_dimensions["A"].width = 120
    wb.move_sheet(how, offset=-1)
    wb.active = 0
    wb.save(path)
    done = sum(1 for r in rows.values() if r["langs"].get(lang))
    print(f"{path}: {len(rows)} texts, {done} of them in {lang} already")


def import_(lang, path):
    wb = load_workbook(path, read_only=True)
    ws = wb["Texts"]
    rows = ws.iter_rows(values_only=True)
    head = [str(h or "") for h in next(rows)]
    col = next(i for i, h in enumerate(head) if h.endswith("(" + lang + ")"))
    got, bad = {}, []
    for row in rows:
        key, text = row[1], row[col]
        if not key or text is None or not str(text).strip():
            continue
        key, text = str(key).replace("\r\n", "\n"), str(text).replace("\r\n", "\n").strip()
        if sorted(PLACEHOLDER.findall(key)) != sorted(PLACEHOLDER.findall(text)):
            bad.append((row[0], key, text))
            continue
        got[key] = text

    changed = 0
    for table, lines, eol in tables():
        out = list(lines)
        # (from the bottom, so the line numbers above stay right)
        for key, at, langs in reversed(entries(lines)):
            text = got.get(key)
            if text is None:
                continue
            line = f"  {lang}: {quote(text)}"
            if lang in langs:
                out[langs[lang][0]] = line
            else:
                last = max([at] + [i for i, _ in langs.values()])
                out.insert(last + 1, line)
        if out != lines:
            with open(table, "w", encoding="utf-8", newline="") as f:
                f.write(eol.join(out))
            changed += 1
    print(f"{len(got)} texts in {lang}, {changed} tables written")
    for n, key, text in bad:
        print(f"  #{n}: the placeholders differ from the English text's, left out\n    {key!r}\n    {text!r}", file=sys.stderr)


def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = p.add_subparsers(dest="cmd", required=True)
    e = sub.add_parser("export")
    e.add_argument("lang")
    e.add_argument("--ref", default="es,fr", help="the languages shown beside it (default es,fr)")
    e.add_argument("--out")
    i = sub.add_parser("import")
    i.add_argument("lang")
    i.add_argument("file")
    a = p.parse_args()
    lang = a.lang.lower()
    if a.cmd == "export":
        export(lang, [r for r in a.ref.lower().split(",") if r and r != lang], a.out or f"omsihub-{lang}.xlsx")
    else:
        import_(lang, a.file)


if __name__ == "__main__":
    main()
